"""
 * @file: catalog.py
 * @description: Роутер для нового каталога товаров
 * @dependencies: models, services, templates
 * @created: 2024-12-19
"""

from typing import List, Optional
from fastapi import APIRouter, Depends, Query, HTTPException, Request
from sqlmodel import select, func
from sqlmodel.ext.asyncio.session import AsyncSession
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse, StreamingResponse
from sqlalchemy import or_, and_
from fastapi.templating import Jinja2Templates
from app.api import deps
from app.models.warehouse import Product, Stock
from app.models.user import User
from app.services.warehouse.manager import Warehouses
from app.services.prices_service import prices_service
import logging
import io
import csv
from datetime import datetime
import base64
import pandas as pd
from io import BytesIO

logger = logging.getLogger(__name__)

router = APIRouter()
templates = Jinja2Templates(directory="app/templates")

@router.get("/", response_class=HTMLResponse)
async def catalog_page(
    request: Request,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=1000),
    search: Optional[str] = None,
    sku_search: Optional[str] = None,
    stock_filter: Optional[int] = None,
    min_stock_filter: Optional[int] = None,
    brand_filter: Optional[str] = None,
    sort_order: Optional[str] = None,
    operation_skus: Optional[str] = None,
    db: AsyncSession = Depends(deps.get_async_session),
    current_user: User = Depends(deps.get_current_user_optional)
):
    """
    Главная страница нового каталога товаров
    """
    if not current_user:
        return RedirectResponse(url=f"/login?next=/catalog_new", status_code=302)

    # Логирование параметров запроса
    logger.info(f"New catalog request: search={search}, sku_search={sku_search}, stock_filter={stock_filter}, min_stock_filter={min_stock_filter}, brand_filter={brand_filter}")

    # Валидация и автоматическая коррекция фильтров
    if (stock_filter is not None and min_stock_filter is not None and 
        stock_filter < min_stock_filter):
        stock_filter, min_stock_filter = min_stock_filter, stock_filter

    # Базовый запрос для товаров с подсчетом остатков
    base_query = (
        select(
            Product,
            func.sum(Stock.quantity).label('total_stock')
        )
        .outerjoin(Stock, Stock.sku == Product.sku)
        .group_by(Product.sku)
    )
    
    # Поиск по названию, SKU или EAN
    if search:
        name_sku_search = f"%{search}%"
        base_query = base_query.where(
            or_(
                Product.name.ilike(name_sku_search),
                Product.sku.ilike(name_sku_search),
                Product.eans.contains([search])
            )
        )

    # Поиск по списку SKU (поддержка запятых и пробелов)
    if sku_search:
        # Разделяем по запятым и пробелам, убираем пустые элементы
        sku_list = [sku.strip() for sku in sku_search.replace(',', ' ').split() if sku.strip()]
        if sku_list:
            base_query = base_query.where(Product.sku.in_(sku_list))

    # Фильтр по SKU операции (приоритетный фильтр)
    if operation_skus:
        sku_list = [sku.strip() for sku in operation_skus.split(',') if sku.strip()]
        if sku_list:
            base_query = base_query.where(Product.sku.in_(sku_list))

    # Собираем условия фильтров по количеству
    having_conditions = []

    # Фильтр по максимальному количеству
    if stock_filter is not None:
        having_conditions.append(func.coalesce(func.sum(Stock.quantity), 0) < stock_filter)

    # Фильтр по минимальному количеству
    if min_stock_filter is not None:
        having_conditions.append(func.coalesce(func.sum(Stock.quantity), 0) >= min_stock_filter)

    # Применяем условия, если они есть
    if having_conditions:
        base_query = base_query.having(and_(*having_conditions))

    # Фильтр по бренду
    if brand_filter:
        if brand_filter == "__no_brand__":
            # Фильтр по товарам без бренда (NULL или пустая строка)
            base_query = base_query.where(
                or_(
                    Product.brand.is_(None),
                    Product.brand == "",
                    Product.brand == " "
                )
            )
        else:
            base_query = base_query.where(Product.brand == brand_filter)

    # Создаем подзапрос для корректного подсчета общего количества
    subquery = base_query.subquery()
    
    # Получаем общее количество товаров для пагинации
    total_count_query = select(func.count()).select_from(subquery)
    result = await db.exec(total_count_query)
    total_count = result.first()
    total_pages = (total_count + page_size - 1) // page_size

    # Основной запрос с сортировкой
    query = (
        select(Product, subquery.c.total_stock)
        .join(subquery, Product.sku == subquery.c.sku)
    )

    # Применяем сортировку
    if sort_order == 'desc':
        query = query.order_by(func.coalesce(subquery.c.total_stock, 0).desc(), Product.sku.asc())
    else:
        # По умолчанию и для sort_order == 'asc' сортируем по возрастанию остатков
        query = query.order_by(func.coalesce(subquery.c.total_stock, 0).asc(), Product.sku.asc())

    # Применяем пагинацию
    query = query.offset((page - 1) * page_size).limit(page_size)
    
    # Получаем товары
    result = await db.exec(query)
    products = result.all()

    # Создаем список продуктов с их остатками
    products_with_stocks = []
    for row in products:
        product = row[0]  # Получаем объект Product
        total_stock = row[1] or 0  # Получаем total_stock
        
        # Получаем остатки для продукта
        stocks_query = select(Stock).where(Stock.sku == product.sku)
        result = await db.exec(stocks_query)
        stocks = result.all()
        
        # Создаем словарь с данными продукта
        product_data = {
            "id": product.sku,
            "sku": product.sku,
            "name": product.name,
            "name_eng": product.name_eng,  # Добавляем английское название
            "brand": product.brand,
            "eans": product.eans,
            "ean": product.eans[0] if product.eans else None,
            "image": base64.b64encode(product.image).decode('utf-8') if product.image else None,
            "has_image": bool(product.original_image),  # Проверяем наличие оригинального изображения
            "total_stock": total_stock,
            "stocks": {}
        }
        
        # Инициализируем остатки для всех складов как 0
        for warehouse in Warehouses:
            product_data["stocks"][warehouse.value] = 0
            
        # Заполняем фактические остатки
        for stock in stocks:
            product_data["stocks"][stock.warehouse] = stock.quantity
        
        products_with_stocks.append(product_data)

    # Обогащаем данные о товарах ценами из удаленной БД
    if prices_service.is_available() and products_with_stocks:
        try:
            # Получаем все SKU товаров
            skus = [product["sku"] for product in products_with_stocks]
            
            # Получаем цены для всех SKU одним запросом
            prices_data = prices_service.get_prices_by_skus(skus)
            
            # Добавляем цены к данным товаров
            for product_data in products_with_stocks:
                sku = product_data["sku"]
                price_info = prices_data.get(sku)
                if price_info and price_info.min_price:
                    # Конвертируем Decimal в float для JSON сериализации
                    product_data["min_price"] = float(price_info.min_price)
                else:
                    product_data["min_price"] = None
        except Exception as e:
            logger.error(f"Error getting prices: {str(e)}")
            # Если не удалось получить цены, добавляем пустые значения
            for product_data in products_with_stocks:
                product_data["min_price"] = None

    # Для обычных запросов возвращаем HTML-страницу
    warehouses = [w.value for w in Warehouses]
    
    return templates.TemplateResponse(
        "catalog_new.html",
        {
            "request": request,
            "user": current_user,
            "products": products_with_stocks,
            "page": page,
            "page_size": page_size,
            "total_pages": total_pages,
            "warehouses": warehouses,
            "current_user": current_user,
            "filters": {
                "search": search or "",
                "stock_filter": stock_filter,
                "min_stock_filter": min_stock_filter,
                "brand_filter": brand_filter or "",
                "sort_order": sort_order or "",
                "page_size": page_size
            }
        }
    )

@router.get("/api/products")
async def get_products_api(
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=1000),
    search: Optional[str] = None,
    sku_search: Optional[str] = None,
    stock_filter: Optional[int] = None,
    min_stock_filter: Optional[int] = None,
    brand_filter: Optional[str] = None,
    sort_order: Optional[str] = None,
    operation_skus: Optional[str] = None,
    db: AsyncSession = Depends(deps.get_async_session),
    current_user: User = Depends(deps.get_current_user_optional)
):
    """
    API эндпоинт для получения товаров в формате JSON
    """
    if not current_user:
        raise HTTPException(status_code=401, detail="Требуется авторизация")

    # Логирование параметров запроса
    logger.info(f"Catalog API request: search={search}, sku_search={sku_search}, stock_filter={stock_filter}, min_stock_filter={min_stock_filter}, brand_filter={brand_filter}")

    # Валидация и автоматическая коррекция фильтров
    if (stock_filter is not None and min_stock_filter is not None and 
        stock_filter < min_stock_filter):
        stock_filter, min_stock_filter = min_stock_filter, stock_filter

    # Базовый запрос для товаров с подсчетом остатков
    base_query = (
        select(
            Product,
            func.sum(Stock.quantity).label('total_stock')
        )
        .outerjoin(Stock, Stock.sku == Product.sku)
        .group_by(Product.sku)
    )
    
    # Поиск по названию, SKU или EAN
    if search:
        name_sku_search = f"%{search}%"
        base_query = base_query.where(
            or_(
                Product.name.ilike(name_sku_search),
                Product.sku.ilike(name_sku_search),
                Product.eans.contains([search])
            )
        )

    # Поиск по списку SKU (поддержка запятых и пробелов)
    if sku_search:
        # Разделяем по запятым и пробелам, убираем пустые элементы
        sku_list = [sku.strip() for sku in sku_search.replace(',', ' ').split() if sku.strip()]
        if sku_list:
            base_query = base_query.where(Product.sku.in_(sku_list))

    # Фильтр по SKU операции (приоритетный фильтр)
    if operation_skus:
        sku_list = [sku.strip() for sku in operation_skus.split(',') if sku.strip()]
        if sku_list:
            base_query = base_query.where(Product.sku.in_(sku_list))

    # Собираем условия фильтров по количеству
    having_conditions = []

    # Фильтр по максимальному количеству
    if stock_filter is not None:
        having_conditions.append(func.coalesce(func.sum(Stock.quantity), 0) < stock_filter)

    # Фильтр по минимальному количеству
    if min_stock_filter is not None:
        having_conditions.append(func.coalesce(func.sum(Stock.quantity), 0) >= min_stock_filter)

    # Применяем условия, если они есть
    if having_conditions:
        base_query = base_query.having(and_(*having_conditions))

    # Фильтр по бренду
    if brand_filter:
        if brand_filter == "__no_brand__":
            # Фильтр по товарам без бренда (NULL или пустая строка)
            base_query = base_query.where(
                or_(
                    Product.brand.is_(None),
                    Product.brand == "",
                    Product.brand == " "
                )
            )
        else:
            base_query = base_query.where(Product.brand == brand_filter)

    # Создаем подзапрос для корректного подсчета общего количества
    subquery = base_query.subquery()
    
    # Получаем общее количество товаров для пагинации
    total_count_query = select(func.count()).select_from(subquery)
    result = await db.exec(total_count_query)
    total_count = result.first()
    total_pages = (total_count + page_size - 1) // page_size

    # Основной запрос с сортировкой
    query = (
        select(Product, subquery.c.total_stock)
        .join(subquery, Product.sku == subquery.c.sku)
    )

    # Применяем сортировку
    if sort_order == 'desc':
        query = query.order_by(func.coalesce(subquery.c.total_stock, 0).desc(), Product.sku.asc())
    else:
        # По умолчанию и для sort_order == 'asc' сортируем по возрастанию остатков
        query = query.order_by(func.coalesce(subquery.c.total_stock, 0).asc(), Product.sku.asc())

    # Применяем пагинацию
    query = query.offset((page - 1) * page_size).limit(page_size)
    
    # Получаем товары
    result = await db.exec(query)
    products = result.all()

    # Создаем список продуктов с их остатками
    products_with_stocks = []
    for row in products:
        product = row[0]  # Получаем объект Product
        total_stock = row[1] or 0  # Получаем total_stock
        
        # Получаем остатки для продукта
        stocks_query = select(Stock).where(Stock.sku == product.sku)
        result = await db.exec(stocks_query)
        stocks = result.all()
        
        # Создаем словарь с данными продукта
        product_data = {
            "id": product.sku,
            "sku": product.sku,
            "name": product.name,
            "name_eng": product.name_eng,  # Добавляем английское название
            "brand": product.brand,
            "eans": product.eans,
            "ean": product.eans[0] if product.eans else None,
            "image": base64.b64encode(product.image).decode('utf-8') if product.image else None,
            "has_image": bool(product.original_image),  # Проверяем наличие оригинального изображения
            "total_stock": total_stock,
            "stocks": {}
        }
        
        # Инициализируем остатки для всех складов как 0
        for warehouse in Warehouses:
            product_data["stocks"][warehouse.value] = 0
            
        # Заполняем фактические остатки
        for stock in stocks:
            product_data["stocks"][stock.warehouse] = stock.quantity
        
        products_with_stocks.append(product_data)

    # Обогащаем данные о товарах ценами из удаленной БД
    if prices_service.is_available() and products_with_stocks:
        try:
            # Получаем все SKU товаров
            skus = [product["sku"] for product in products_with_stocks]
            
            # Получаем цены для всех SKU одним запросом
            prices_data = prices_service.get_prices_by_skus(skus)
            
            # Добавляем цены к данным товаров
            for product_data in products_with_stocks:
                sku = product_data["sku"]
                price_info = prices_data.get(sku)
                if price_info and price_info.min_price:
                    # Конвертируем Decimal в float для JSON сериализации
                    product_data["min_price"] = float(price_info.min_price)
                else:
                    product_data["min_price"] = None
        except Exception as e:
            logger.error(f"Error getting prices: {str(e)}")
            # Если не удалось получить цены, добавляем пустые значения
            for product_data in products_with_stocks:
                product_data["min_price"] = None

    # Для AJAX-запросов генерируем HTML для каждой карточки
    products_response = []
    for product_data in products_with_stocks:
        html = templates.get_template("catalog_new_blocks/product_card.html").render(
            product=product_data,
            selected_products=[],
            current_user=current_user
        )
        products_response.append({
            "data": product_data,
            "html": html
        })

    return {
        "products": products_response,
        "total": total_count,
        "page": page,
        "page_size": page_size,
        "total_pages": total_pages
    }

@router.get("/api/products/{sku}/card")
async def get_product_card_new(
    request: Request,
    sku: str,
    db: AsyncSession = Depends(deps.get_async_session),
    current_user: User = Depends(deps.get_current_user_optional)
):
    """
    Получить HTML карточки товара по SKU для нового каталога
    """
    if not current_user:
        raise HTTPException(status_code=401, detail="Необходима авторизация")
    
    # Получаем товар с остатками
    product_query = (
        select(
            Product,
            func.sum(Stock.quantity).label('total_stock')
        )
        .outerjoin(Stock, Stock.sku == Product.sku)
        .where(Product.sku == sku)
        .group_by(Product.sku)
    )
    
    result = await db.exec(product_query)
    product_row = result.first()
    
    if not product_row:
        raise HTTPException(status_code=404, detail="Товар не найден")
    
    product = product_row[0]
    total_stock = product_row[1] or 0
    
    # Получаем остатки для продукта
    stocks_query = select(Stock).where(Stock.sku == product.sku)
    result = await db.exec(stocks_query)
    stocks = result.all()
    
    # Создаем словарь с данными продукта
    product_data = {
        "id": product.sku,
        "sku": product.sku,
        "name": product.name,
        "brand": product.brand,
        "eans": product.eans,
        "ean": product.eans[0] if product.eans else None,
        "image": base64.b64encode(product.image).decode('utf-8') if product.image else None,
        "has_image": bool(product.original_image),  # Проверяем наличие оригинального изображения
        "total_stock": total_stock,
        "stocks": {}
    }
    
    # Инициализируем остатки для всех складов как 0
    for warehouse in Warehouses:
        product_data["stocks"][warehouse.value] = 0
        
    # Заполняем фактические остатки
    for stock in stocks:
        product_data["stocks"][stock.warehouse] = stock.quantity
    
    # Получаем цену если доступен сервис цен
    if prices_service.is_available():
        try:
            price_info = prices_service.get_price_by_sku(product.sku)
            if price_info and price_info.min_price:
                # Конвертируем Decimal в float для JSON сериализации
                product_data["min_price"] = float(price_info.min_price)
            else:
                product_data["min_price"] = None
        except Exception as e:
            logger.error(f"Ошибка получения цены для SKU {sku}: {str(e)}")
            product_data["min_price"] = None
    else:
        product_data["min_price"] = None
    
    # Генерируем HTML карточки товара
    html = templates.get_template("catalog_new_blocks/product_card.html").render(
        product=product_data,
        selected_products=[],
        current_user=current_user
    )
    
    return JSONResponse(content={
        "html": html,
        "data": product_data
    })

@router.get("/api/brands")
async def get_brands_api(
    db: AsyncSession = Depends(deps.get_async_session),
    current_user: User = Depends(deps.get_current_user_optional)
):
    """
    API эндпоинт для получения списка брендов
    """
    if not current_user:
        raise HTTPException(status_code=401, detail="Требуется авторизация")
    
    try:
        # Получаем уникальные бренды, исключая NULL значения
        query = select(Product.brand).where(Product.brand.is_not(None)).distinct().order_by(Product.brand)
        result = await db.exec(query)
        brands = result.all()
        
        # Возвращаем список брендов
        return {"brands": [brand for brand in brands if brand and brand.strip()]}
    except Exception as e:
        logger.error(f"Ошибка при получении списка брендов: {str(e)}")
        raise HTTPException(
            status_code=500,
            detail=f"Не удалось получить список брендов: {str(e)}"
        )

# ==================== WORKSPACE ACTIONS API ====================


@router.post("/api/workspace/sync/enable")
async def enable_sync_for_workspace(
    request: Request,
    skus: List[str],
    db: AsyncSession = Depends(deps.get_async_session),
    current_user: User = Depends(deps.get_current_user_optional)
):
    """
    Включить синхронизацию остатков для выбранных товаров
    """
    if not current_user:
        raise HTTPException(status_code=401, detail="Требуется авторизация")
    
    if not skus:
        raise HTTPException(status_code=400, detail="Список SKU не может быть пустым")
    
    try:
        # Получаем токены из микросервиса Allegro
        from app.core.security import create_access_token
        from app.core.config import settings
        from app.services.Allegro_Microservice.tokens_endpoint import AllegroTokenMicroserviceClient
        from app.services.product_allegro_sync_service import ProductAllegroSyncService
        
        # Создаем JWT токен для аутентификации с микросервисом
        jwt_token = create_access_token(user_id=settings.PROJECT_NAME)
        
        # Получаем активные токены из микросервиса
        token_client = AllegroTokenMicroserviceClient(
            jwt_token=jwt_token
        )
        
        tokens_response = token_client.get_tokens(per_page=100, active_only=True)
        
        if not tokens_response.items:
            raise HTTPException(status_code=404, detail="Активные токены Allegro не найдены")
        
        # Создаем сервис синхронизации
        sync_service = ProductAllegroSyncService(db)
        
        # Для каждого токена создаем или обновляем настройки синхронизации
        total_updated = 0
        accounts_processed = []
        
        for token_data in tokens_response.items:
            account_name = token_data['account_name']
            accounts_processed.append(account_name)
            
            # Для каждого SKU создаем или обновляем настройки синхронизации
            for sku in skus:
                try:
                    await sync_service.create_or_update_sync_settings(
                        product_sku=sku,
                        account_name=account_name,
                        stock_sync_enabled=True,
                        price_sync_enabled=False,
                        price_multiplier=1.0
                    )
                    total_updated += 1
                except Exception as e:
                    logger.error(f"Ошибка при создании настроек синхронизации для SKU {sku} и аккаунта {account_name}: {str(e)}")
        
        logger.info(f"Включена синхронизация для {total_updated} товар-аккаунт пар из {len(skus)} товаров и {len(accounts_processed)} аккаунтов")
        
        return {
            "message": f"Синхронизация включена для {total_updated} товар-аккаунт пар",
            "updated_count": total_updated,
            "requested_skus": len(skus),
            "accounts_processed": len(accounts_processed),
            "accounts": accounts_processed
        }
        
    except Exception as e:
        logger.error(f"Ошибка при включении синхронизации: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Ошибка включения синхронизации: {str(e)}")

@router.post("/api/workspace/sync/disable")
async def disable_sync_for_workspace(
    request: Request,
    skus: List[str],
    db: AsyncSession = Depends(deps.get_async_session),
    current_user: User = Depends(deps.get_current_user_optional)
):
    """
    Выключить синхронизацию остатков для выбранных товаров
    """
    if not current_user:
        raise HTTPException(status_code=401, detail="Требуется авторизация")
    
    if not skus:
        raise HTTPException(status_code=400, detail="Список SKU не может быть пустым")
    
    try:
        # Получаем токены из микросервиса Allegro
        from app.core.security import create_access_token
        from app.core.config import settings
        from app.services.Allegro_Microservice.tokens_endpoint import AllegroTokenMicroserviceClient
        from app.services.product_allegro_sync_service import ProductAllegroSyncService
        
        # Создаем JWT токен для аутентификации с микросервисом
        jwt_token = create_access_token(user_id=settings.PROJECT_NAME)
        
        # Получаем активные токены из микросервиса
        token_client = AllegroTokenMicroserviceClient(
            jwt_token=jwt_token
        )
        
        tokens_response = token_client.get_tokens(per_page=100, active_only=True)
        
        if not tokens_response.items:
            raise HTTPException(status_code=404, detail="Активные токены Allegro не найдены")
        
        # Создаем сервис синхронизации
        sync_service = ProductAllegroSyncService(db)
        
        # Для каждого токена отключаем настройки синхронизации
        total_updated = 0
        accounts_processed = []
        
        for token_data in tokens_response.items:
            account_name = token_data['account_name']
            accounts_processed.append(account_name)
            
            # Для каждого SKU отключаем настройки синхронизации
            for sku in skus:
                try:
                    await sync_service.create_or_update_sync_settings(
                        product_sku=sku,
                        account_name=account_name,
                        stock_sync_enabled=False,
                        price_sync_enabled=False,
                        price_multiplier=1.0
                    )
                    total_updated += 1
                except Exception as e:
                    logger.error(f"Ошибка при отключении настроек синхронизации для SKU {sku} и аккаунта {account_name}: {str(e)}")
        
        logger.info(f"Отключена синхронизация для {total_updated} товар-аккаунт пар из {len(skus)} товаров и {len(accounts_processed)} аккаунтов")
        
        return {
            "message": f"Синхронизация отключена для {total_updated} товар-аккаунт пар",
            "updated_count": total_updated,
            "requested_skus": len(skus),
            "accounts_processed": len(accounts_processed),
            "accounts": accounts_processed
        }
        
    except Exception as e:
        logger.error(f"Ошибка при отключении синхронизации: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Ошибка отключения синхронизации: {str(e)}")

@router.post("/api/workspace/transfer-request")
async def create_transfer_request(
    request: Request,
    skus: List[str],
    db: AsyncSession = Depends(deps.get_async_session),
    current_user: User = Depends(deps.get_current_user_optional)
):
    """
    Создать файл заявки на перемещение для выбранных товаров
    """
    if not current_user:
        raise HTTPException(status_code=401, detail="Требуется авторизация")
    
    if not skus:
        raise HTTPException(status_code=400, detail="Список SKU не может быть пустым")
    
    try:
        # Получаем товары с остатками
        products_query = (
            select(
                Product,
                func.sum(Stock.quantity).label('total_stock')
            )
            .outerjoin(Stock, Stock.sku == Product.sku)
            .where(Product.sku.in_(skus))
            .group_by(Product.sku)
        )
        
        result = await db.exec(products_query)
        products = result.all()
        
        if not products:
            raise HTTPException(status_code=404, detail="Товары не найдены")
        
        # Создаем CSV файл заявки на перемещение
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Заголовки для заявки на перемещение
        writer.writerow([
            'SKU', 'Название', 'Бренд', 'EAN', 'Общий остаток',
            'Склад 1', 'Склад 2', 'Склад 3', 'Склад 4', 'Склад 5',
            'Количество для перемещения', 'Примечание'
        ])
        
        # Данные товаров
        for product_row in products:
            product = product_row[0]
            total_stock = product_row[1] or 0
            
            # Получаем остатки по складам
            stocks_query = select(Stock).where(Stock.sku == product.sku)
            stocks_result = await db.exec(stocks_query)
            stocks = stocks_result.all()
            
            # Создаем словарь остатков по складам
            stocks_dict = {}
            for stock in stocks:
                stocks_dict[stock.warehouse] = stock.quantity
            
            # Формируем строку данных
            row_data = [
                product.sku,
                product.name or '',
                product.brand or '',
                product.eans[0] if product.eans else '',
                total_stock,
                stocks_dict.get('warehouse_1', 0),
                stocks_dict.get('warehouse_2', 0),
                stocks_dict.get('warehouse_3', 0),
                stocks_dict.get('warehouse_4', 0),
                stocks_dict.get('warehouse_5', 0),
                '',  # Количество для перемещения (заполняется вручную)
                ''   # Примечание (заполняется вручную)
            ]
            
            writer.writerow(row_data)
        
        # Подготавливаем файл для скачивания
        output.seek(0)
        csv_content = output.getvalue()
        output.close()
        
        # Создаем имя файла с датой
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transfer_request_{timestamp}.csv"
        
        return StreamingResponse(
            io.BytesIO(csv_content.encode('utf-8')),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
        
    except Exception as e:
        logger.error(f"Ошибка при создании заявки на перемещение: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Ошибка создания заявки: {str(e)}")

@router.post("/api/workspace/supplier-request")
async def create_supplier_request(
    request: Request,
    products_data: List[dict],
    db: AsyncSession = Depends(deps.get_async_session),
    current_user: User = Depends(deps.get_current_user_optional)
):
    """
    Создать Excel файл заявки поставщику для выбранных товаров
    """
    if not current_user:
        raise HTTPException(status_code=401, detail="Требуется авторизация")
    
    if not products_data:
        raise HTTPException(status_code=400, detail="Список товаров не может быть пустым")
    
    try:
        # Извлекаем SKU из данных
        skus = [item['sku'] for item in products_data]
        
        # Получаем товары с остатками
        products_query = (
            select(
                Product,
                func.sum(Stock.quantity).label('total_stock')
            )
            .outerjoin(Stock, Stock.sku == Product.sku)
            .where(Product.sku.in_(skus))
            .group_by(Product.sku)
        )
        
        result = await db.exec(products_query)
        products = result.all()
        
        if not products:
            raise HTTPException(status_code=404, detail="Товары не найдены")
        
        # Создаем словарь для быстрого поиска данных товаров
        products_dict = {item['sku']: item for item in products_data}
        
        # Создаем DataFrame для удобной работы
        import pandas as pd
        from openpyxl.utils import get_column_letter
        from openpyxl.drawing.image import Image as XLImage
        
        data = []
        images = []
        
        for product_row in products:
            product = product_row[0]
            product_data = products_dict.get(product.sku, {})
            
            row_data = {
                'No': len(data) + 1,
                'Brand': product.brand or '',
                'Foto': '',  # Пустое поле для изображения
                'Name': product.name_eng or product.name or '',
                'EAN/UPC': product.eans[0] if product.eans else '',
                'SKU': product.sku,
                'Количество': product_data.get('quantity', 0),
                'Комментарий': product_data.get('comment', '')
            }
            data.append(row_data)
            
            # Добавляем изображение если есть
            if product.image:
                images.append(product.image)
            else:
                images.append(None)
        
        df = pd.DataFrame(data)
        
        # Создаем Excel файл
        excel_buffer = BytesIO()
        with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
            # Пишем данные (картинки пока пустые)
            df.to_excel(writer, index=False, sheet_name="Заявка поставщику")
            ws = writer.sheets["Заявка поставщику"]
            
            # Находим букву колонки «Foto»
            foto_col_idx = df.columns.get_loc("Foto") + 1
            foto_col_letter = get_column_letter(foto_col_idx)
            
            # Форматируем заголовки и ширину колонок
            from openpyxl.styles import Alignment
            
            for idx, col in enumerate(df.columns, start=1):
                cell = ws.cell(row=1, column=idx)
                cell.font = cell.font.copy(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
                if col == "Foto":
                    ws.column_dimensions[foto_col_letter].width = 22  # ~150px
                elif col == "Name":
                    ws.column_dimensions[get_column_letter(idx)].width = 100  # Для длинных названий
                elif col == "Комментарий":
                    ws.column_dimensions[get_column_letter(idx)].width = 50  # Для комментариев
                else:
                    max_len = max(df[col].astype(str).map(len).max(), len(col))
                    ws.column_dimensions[get_column_letter(idx)].width = min(max_len + 2, 50)
            
            # Центрируем содержимое всех ячеек с данными
            for row_idx in range(2, len(df) + 2):  # Начинаем с 2-й строки (данные)
                for col_idx in range(1, len(df.columns) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Вставка картинок 150×150 и высота строк
            for row_idx, img_bytes in enumerate(images, start=2):
                if not img_bytes:
                    continue
                try:
                    img = XLImage(BytesIO(img_bytes))
                    img.width = 150
                    img.height = 150
                    ws.row_dimensions[row_idx].height = 115  # под ~150px
                    
                    ws.add_image(img, f"{foto_col_letter}{row_idx}")
                except Exception as e:
                    logger.error(f"Не удалось вставить изображение в строке {row_idx}: {e}")
        
        excel_buffer.seek(0)
        
        # Создаем имя файла с датой
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"supplier_request_{timestamp}.xlsx"
        
        return StreamingResponse(
            excel_buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
        
    except Exception as e:
        logger.error(f"Ошибка при создании заявки поставщику: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Ошибка создания заявки: {str(e)}")
